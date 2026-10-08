"""
数式の機械判定スクリプト。
AIが返した「ある1行の数式」を、同じ区分の全行へ"下にコピーした"ときの
相対参照のズレを再現してから、実際にopenpyxl+LibreOfficeで再計算し、
真値（Pythonで直接集計した値）と比較する。
"""
import re, json, subprocess, sys, os
import openpyxl

CELL_RE = re.compile(r'(?<![A-Za-z0-9_$])(\$?)([A-Z]{1,3})(\$?)(\d+)')

def shift_formula(formula, row_delta):
    """数式中の相対行参照だけを row_delta だけ平行移動する（$固定は動かさない）。"""
    def repl(m):
        col_abs, col, row_abs, row = m.groups()
        row_n = int(row)
        if row_abs == '$':
            new_row = row_n
        else:
            new_row = row_n + row_delta
        return f"{col_abs}{col}{row_abs}{new_row}"
    return CELL_RE.sub(repl, formula)

def build_and_check(sheet_json_path, formulas_by_target, out_xlsx):
    """
    sheet_json_path: build_sheet.py が作った *_sheet.json
    formulas_by_target: {"store": (example_row, "=formula"), "area": (...), ...}
      example_row = そのtargetの最初の出現行番号、formula = その行に実際に入れる式
    """
    d = json.load(open(sheet_json_path, encoding="utf-8"))
    sheet, gt = d["sheet"], {int(k): v for k, v in d["gt"].items()}

    wb = openpyxl.Workbook()
    ws = wb.active
    cols = ["A", "B", "C", "D", "E", "F"]
    header = list(sheet[0].keys())
    # header row 1
    headerslabels = ["A", "B", "C", "D", "E", "F"]
    for i, c in enumerate(cols):
        ws[f"{c}1"] = headerslabels[i]

    target_rows = {}  # target -> list of row numbers
    for rec in sheet:
        r = rec["row"]
        ws[f"A{r}"] = rec["A"]
        ws[f"B{r}"] = rec["B"]
        ws[f"C{r}"] = rec["C"]
        ws[f"D{r}"] = rec["D"]
        ws[f"E{r}"] = rec["E"]
        if rec["target"] is None:
            ws[f"F{r}"] = rec["F"]
        else:
            target_rows.setdefault(rec["target"], []).append(r)

    applied = {}  # row -> formula actually written
    missing_targets = []
    for target, rows in target_rows.items():
        if target not in formulas_by_target:
            missing_targets.append(target)
            continue
        example_row, formula = formulas_by_target[target]
        for r in rows:
            f2 = shift_formula(formula, r - example_row)
            ws[f"F{r}"] = f2
            applied[r] = f2

    wb.save(out_xlsx)

    # LibreOffice headless recalc -> convert to csv forces full recalculation
    out_dir = os.path.dirname(os.path.abspath(out_xlsx))
    subprocess.run(
        ["soffice", "--headless", "--calc", "--convert-to",
         "csv:Text - txt - csv (StarCalc):44,34,0,1,,0,false,true,false,false,false",
         "--outdir", out_dir, out_xlsx],
        check=True, capture_output=True, timeout=120
    )
    csv_path = os.path.join(out_dir, os.path.splitext(os.path.basename(out_xlsx))[0] + ".csv")
    import csv as csvmod
    rows_out = list(csvmod.reader(open(csv_path, encoding="utf-8")))

    results = {}
    for target, rows in target_rows.items():
        for r in rows:
            if r > len(rows_out):
                results[r] = ("NO_ROW", None, gt.get(r))
                continue
            raw = rows_out[r - 1][5] if len(rows_out[r - 1]) > 5 else ""
            try:
                val = int(round(float(raw.replace(",", ""))))
            except Exception:
                val = raw
            truth = gt.get(r)
            ok = (val == truth)
            results[r] = (ok, val, truth)
    return dict(results=results, missing_targets=missing_targets, target_rows=target_rows, applied_example={t: formulas_by_target[t] for t in formulas_by_target})

if __name__ == "__main__":
    pass
